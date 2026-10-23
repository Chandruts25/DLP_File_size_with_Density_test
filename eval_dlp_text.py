"""Evaluate the DLP /api/m/dlp/text/ endpoint on a folder of sample files.

Sends each chosen .txt file (multipart field `data`, mode=full by default), times it,
compares what the DLP found with the folder's answer key, and writes a two-sheet Excel
report: Summary (headline figures) and Results (time and accuracy per file).

Pick the folder (a name next to this script, or any path):
    python eval_dlp_text.py --dir dlp_samples_low                    # every file, once
Pick how many times each file is sent:
    python eval_dlp_text.py --dir dlp_samples_low --runs 3
Pick which files (numbers are the file's position in size order, as # in the report):
    python eval_dlp_text.py --dir dlp_samples_high --only 27         # just file 27
    python eval_dlp_text.py --dir dlp_samples_high --only 5 --only 9 # files 5 and 9
    python eval_dlp_text.py --dir dlp_samples_high --first 10        # files 1-10
    python eval_dlp_text.py --dir dlp_samples_high --from 20 --to 30 # files 20-30
Limits (both configurable):
    --timeout 600              give up on one request after 600 s (default 10 min)
    --stop-after-timeouts 2    stop the whole run after 2 timeouts in a row (default 2)

Files are sent one at a time, smallest first. Every run is appended to a .jsonl log as
soon as it finishes and the report is rebuilt after each file, so an interrupted run
still leaves a usable report. Running the same command again resumes: runs already in
the log are skipped (--retry-failed also re-sends files whose runs all failed).
--report-only rebuilds the report from the log without sending anything.
"""
import argparse
import csv
import glob
import json
import os
import re
import time
from collections import Counter
from datetime import datetime

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_URL = "https://qaaidlphw.dlpdefence.io/api/m/dlp/text/"
TEXT_TYPES = {"Person name", "Home address", "Disease", "Medication"}   # matched loosely (see score)
FUZZY_LIMIT = 20_000_000                                                # pair checks before loose matching is skipped


# ----------------------------------------------------------------------------
# Answer key and scoring
# ----------------------------------------------------------------------------
def norm(value):
    """Comparison form of a value: lower case, letters and digits only, so
    '4111 1111 1111 1111' and '4111111111111111' are the same card."""
    return re.sub(r"[^0-9a-z]", "", str(value).casefold())


def load_answer_key(folder, name):
    """[(type, value)] planted in `name`, from answer_key/<file>.csv (or its _partN
    pieces), else from answer_key.csv. None when the folder has no answer key."""
    stem = os.path.splitext(name)[0]
    parts = sorted(glob.glob(os.path.join(folder, "answer_key", f"{glob.escape(stem)}.csv")) +
                   glob.glob(os.path.join(folder, "answer_key", f"{glob.escape(stem)}_part*.csv")))
    rows = []
    if parts:
        for part in parts:
            with open(part, encoding="utf-8", newline="") as fh:
                rows += [(r["type"], r["value"]) for r in csv.DictReader(fh)]
        return rows
    combined = os.path.join(folder, "answer_key.csv")
    if os.path.exists(combined):
        with open(combined, encoding="utf-8", newline="") as fh:
            return [(r["type"], r["value"]) for r in csv.DictReader(fh) if r["file"] == name]
    return None


def found_values(analysis):
    """[(group, label, value)] for everything the scan reported under PII / PFI / PHI."""
    out = []
    for group in ("PII", "PFI", "PHI"):
        block = analysis.get(group)
        if not isinstance(block, dict):
            continue
        for label, values in block.items():
            if isinstance(values, dict):
                values = values.get("values", [])
            if isinstance(values, list):
                out += [(group, label, v if isinstance(v, str) else json.dumps(v)) for v in values]
    return out


def score(planted, found):
    """Compare distinct planted values with distinct found values.

    A planted value counts as found when the DLP reported the same value (ignoring
    spaces, dashes and case). For names, addresses, diseases and medications, a
    reported value that contains the planted one, or is contained in it, also counts,
    since the DLP may return "diabetes" for "Type 2 diabetes mellitus".
    """
    p_type = {}
    for kind, value in planted:
        key = norm(value)
        if key and key not in p_type:
            p_type[key] = kind
    f_label = {}
    for group, label, value in found:
        key = norm(value)
        if key and key not in f_label:
            f_label[key] = f"{group}/{label}"

    matched_p = set(p_type) & set(f_label)
    used_f = set(matched_p)
    left_p = [k for k in p_type if k not in matched_p and p_type[k] in TEXT_TYPES]
    left_f = [k for k in f_label if k not in used_f and len(k) >= 4]
    fuzzy_skipped = len(left_p) * len(left_f) > FUZZY_LIMIT
    if not fuzzy_skipped:
        for p in left_p:
            for f in left_f:
                if f in p or p in f:
                    matched_p.add(p)
                    used_f.add(f)
                    break

    by_type = {}
    for key, kind in p_type.items():
        t = by_type.setdefault(kind, [0, 0])
        t[0] += 1
        t[1] += key in matched_p
    planted_n, found_n, matched_n = len(p_type), len(f_label), len(matched_p)
    extra = [k for k in f_label if k not in used_f]
    missed = [k for k in p_type if k not in matched_p]
    recall = matched_n / planted_n if planted_n else None
    precision = (found_n - len(extra)) / found_n if found_n else None
    f1 = (2 * recall * precision / (recall + precision)) if recall and precision else (0.0 if planted_n else None)
    raw_planted = {norm(v): v for _, v in planted}
    raw_found = {norm(v): v for _, _, v in found}
    return {
        "planted": planted_n, "found": found_n, "matched": matched_n, "missed": len(missed), "extra": len(extra),
        "recall": recall, "precision": precision, "f1": f1, "by_type": by_type, "fuzzy_skipped": fuzzy_skipped,
        "missed_examples": [f"{p_type[k]}: {raw_planted[k]}" for k in missed[:10]],
        "extra_examples": [f"{f_label[k]}: {raw_found[k]}" for k in extra[:10]],
    }


# ----------------------------------------------------------------------------
# Calling the API
# ----------------------------------------------------------------------------
def run_once(session, args, path, timeout, planted):
    name = os.path.basename(path)
    headers = {"ClientGUID": "dlp-eval-script"}
    if args.token:
        headers["Authorization"] = f"Bearer {args.token}"
    row = {"file": name, "started_at": datetime.now().isoformat(timespec="seconds")}
    started = time.perf_counter()
    try:
        with open(path, "rb") as fh:
            resp = session.post(args.url, files={"data": (name, fh, "text/plain")}, data={"mode": args.mode},
                                headers=headers, timeout=(60, timeout))
        row["seconds"] = round(time.perf_counter() - started, 2)
        row["http_status"] = resp.status_code
        try:
            body = resp.json()
        except ValueError:
            row["error"] = f"Non-JSON response: {resp.text[:200]}"
            return row
    except requests.exceptions.Timeout:
        row.update(seconds=round(time.perf_counter() - started, 2), http_status=None, error=f"Timed out after {timeout}s")
        return row
    except requests.exceptions.RequestException as exc:
        row.update(seconds=round(time.perf_counter() - started, 2), http_status=None,
                   error=f"{type(exc).__name__}: {exc}"[:300])
        return row

    entry = (body.get("results") or {}).get(name)
    if entry is None and isinstance(body.get("results"), dict) and len(body["results"]) == 1:
        entry = next(iter(body["results"].values()))
    if not isinstance(entry, dict):
        row["error"] = body.get("error") or body.get("detail") or "No result entry for the file"
        return row
    analysis = entry.get("results") or {}
    found = found_values(analysis)
    groups = Counter(g for g, _, _ in found)
    row.update(scan_ok=entry.get("status"), reason=entry.get("reason"), server_seconds=entry.get("computation_time"),
               sensitivity=analysis.get("sensitivity_level"), confidence=analysis.get("confidence_score"),
               pii=groups["PII"], pfi=groups["PFI"], phi=groups["PHI"])
    if entry.get("status") is not True:
        row["error"] = (analysis.get("error") if isinstance(analysis, dict) else None) or entry.get("reason") or "scan failed"
    elif planted is not None:
        row["accuracy"] = score(planted, found)
    return row


# ----------------------------------------------------------------------------
# Excel report
# ----------------------------------------------------------------------------
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
HEAD_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14, color="1F4E78")
OK_FILL, BAD_FILL, WARN_FILL = (PatternFill("solid", fgColor=c) for c in ("E2EFDA", "F8CBAD", "FFF2CC"))
THIN = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
PCT, SEC, MB = "0.0%", "0.0", "0.000"


def _table(ws, cols, rows, widths=None, formats=None, start_row=1, freeze=True):
    for j, c in enumerate(cols, 1):
        cell = ws.cell(start_row, j, c)
        cell.fill, cell.font, cell.border = HEAD_FILL, HEAD_FONT, THIN
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[start_row].height = 32
    for i, row in enumerate(rows, start_row + 1):
        for j, v in enumerate(row, 1):
            cell = ws.cell(i, j, v)
            cell.border = THIN
            fmt = (formats or {}).get(cols[j - 1])
            if fmt:
                cell.number_format = fmt
    for j, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(j)].width = (widths or {}).get(c, 13)
    if freeze:
        ws.freeze_panes = ws.cell(start_row + 1, 3)


def _readable(s):
    return "-" if s is None else (f"{s / 60:.1f} min" if s >= 120 else f"{s:.1f} s")


def _pct(x):
    return None if x is None else round(x, 4)


def summarise(files, runs_by_file, args):
    """One dict per file: timings, status and the accuracy of its last successful run."""
    out = []
    for path, number in files:
        name = os.path.basename(path)
        runs = runs_by_file.get(name, [])
        good = [r for r in runs if r.get("scan_ok") is True and r.get("http_status") == 200]
        secs = [r["seconds"] for r in good]
        server = [r["server_seconds"] for r in good if isinstance(r.get("server_seconds"), (int, float))]
        last = good[-1] if good else (runs[-1] if runs else {})
        acc = next((r["accuracy"] for r in reversed(good) if r.get("accuracy")), None)
        if not runs:
            result = "Not run"
        elif len(good) == len(runs):
            result = "OK"
        elif good:
            result = f"{len(good)}/{len(runs)} OK"
        else:
            result = "Timed out" if all(str(r.get("error", "")).startswith("Timed out") for r in runs) else "FAILED"
        avg = sum(secs) / len(secs) if secs else None
        size_mb = os.path.getsize(path) / 2**20
        out.append({"number": number, "name": name, "size_mb": size_mb, "runs": runs, "good": good, "avg": avg,
                    "min": min(secs) if secs else None, "max": max(secs) if secs else None,
                    "server": sum(server) / len(server) if server else None,
                    "speed": size_mb / (avg / 60) if avg else None, "result": result, "last": last, "acc": acc,
                    "errors": "; ".join(sorted({r.get("error") for r in runs if r.get("error")}))[:300]})
    return out


def build_report(xlsx_path, files, runs_by_file, manifest, args):
    """Two sheets: Summary (the headline figures) and Results (one row per file)."""
    rows = summarise(files, runs_by_file, args)
    all_runs = [r for row in rows for r in row["runs"]]
    good_runs = [r for row in rows for r in row["good"]]
    scored = [row for row in rows if row["acc"]]
    tot = {k: sum(row["acc"][k] for row in scored) for k in ("planted", "found", "matched", "missed", "extra")}
    recall = tot["matched"] / tot["planted"] if tot["planted"] else None
    precision = (tot["found"] - tot["extra"]) / tot["found"] if tot["found"] else None
    f1 = 2 * recall * precision / (recall + precision) if recall and precision else None
    timed = [row["avg"] for row in rows if row["avg"]]

    wb = Workbook()

    # ---- Summary -------------------------------------------------------
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"DLP evaluation - {os.path.basename(os.path.normpath(args.dir))}"
    ws["A1"].font = TITLE_FONT
    facts = [
        ("Files evaluated", len([r for r in rows if r["runs"]])),
        ("Runs per file", args.runs),
        ("Requests succeeded", f"{len(good_runs)} of {len(all_runs)}"),
        ("Total time", _readable(sum(r.get("seconds") or 0 for r in all_runs))),
        ("Average time per file", _readable(sum(timed) / len(timed)) if timed else "-"),
        ("Planted values", tot["planted"] if scored else "-"),
        ("Correctly found", tot["matched"] if scored else "-"),
        ("Missed", tot["missed"] if scored else "-"),
        ("Extra", tot["extra"] if scored else "-"),
        ("Recall", recall), ("Precision", precision), ("F1 score", f1),
    ]
    for i, (k, v) in enumerate(facts, 3):
        ws.cell(i, 1, k).font = Font(bold=True)
        cell = ws.cell(i, 2, v)
        cell.alignment = Alignment(horizontal="left")
        if k in ("Recall", "Precision", "F1 score"):
            cell.number_format = PCT
            if v is not None:
                cell.fill = OK_FILL if v >= 0.9 else WARN_FILL if v >= 0.7 else BAD_FILL

    notes = [
        ("What the columns mean", ""),
        ("Run N (sec)", "Time of that request: upload + scan + response."),
        ("Average (sec)", "Average over the successful runs of the file."),
        ("Result", "OK = every run succeeded; Timed out / FAILED = none did; x/y OK = some did."),
        ("Planted", "Sensitive values put in the file (from the folder's answer key)."),
        ("Found", "Values the DLP reported."),
        ("Correct", "Planted values the DLP reported."),
        ("Missed", "Planted values the DLP did not report."),
        ("Extra", "Values the DLP reported that were not planted."),
        ("Recall", "Correct / Planted: how much of the sensitive data was caught."),
        ("Precision", "Share of the DLP's findings that were really planted."),
        ("F1 score", "Recall and precision combined into one number."),
        ("", ""),
        ("Endpoint", f"{args.url} (mode={args.mode})"),
        ("Limits", f"{args.timeout}s per request; stops after {args.stop_after_timeouts} timeouts in a row."),
        ("Generated", datetime.now().strftime("%d %b %Y %H:%M")),
    ]
    start = len(facts) + 5
    for i, (k, v) in enumerate(notes, start):
        ws.cell(i, 1, k).font = Font(bold=True, color="1F4E78" if k == "What the columns mean" else None)
        ws.cell(i, 2, v)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 80

    # ---- Results -------------------------------------------------------
    res = wb.create_sheet("Results")
    run_cols = [f"Run {i} (sec)" for i in range(1, args.runs + 1)]
    cols = (["#", "File", "Size (MB)"] + run_cols + ["Average (sec)", "Result", "Planted", "Found", "Correct",
                                                     "Missed", "Extra", "Recall", "Precision", "Notes"])
    table = []
    for row in rows:
        a = row["acc"] or {}
        secs = [r.get("seconds") for r in row["runs"]]
        table.append([row["number"], row["name"], round(row["size_mb"], 3)] +
                     [secs[k] if k < len(secs) else None for k in range(args.runs)] +
                     [round(row["avg"], 1) if row["avg"] else None, row["result"], a.get("planted"), a.get("found"),
                      a.get("matched"), a.get("missed"), a.get("extra"), _pct(a.get("recall")), _pct(a.get("precision")),
                      row["errors"]])
    fmts = {c: SEC for c in run_cols + ["Average (sec)"]}
    fmts.update({"Size (MB)": MB, "Recall": PCT, "Precision": PCT})
    _table(res, cols, table, widths={"#": 5, "File": 34, "Result": 11, "Notes": 45}, formats=fmts)
    res_col, rec_col = cols.index("Result") + 1, cols.index("Recall") + 1
    for i, row in enumerate(rows, 2):
        res.cell(i, res_col).fill = (OK_FILL if row["result"] == "OK" else
                                     BAD_FILL if row["result"] in ("FAILED", "Timed out") else WARN_FILL)
        r = row["acc"] and row["acc"].get("recall")
        if r is not None:
            res.cell(i, rec_col).fill = OK_FILL if r >= 0.9 else WARN_FILL if r >= 0.7 else BAD_FILL

    tmp = xlsx_path + ".tmp"
    wb.save(tmp)
    os.replace(tmp, xlsx_path)


# ----------------------------------------------------------------------------
def pick_files(ap, args):
    """[(path, number)] in size order; number = position among all files in the folder."""
    every = sorted((os.path.join(args.dir, f) for f in os.listdir(args.dir) if f.endswith(".txt")), key=os.path.getsize)
    numbered = [(p, i) for i, p in enumerate(every, 1)]
    if args.only:
        chosen = []
        for item in args.only:
            if item.isdigit():
                if not 1 <= int(item) <= len(every):
                    ap.error(f"--only {item}: the folder has {len(every)} files")
                chosen.append(numbered[int(item) - 1])
            else:
                hit = [x for x in numbered if os.path.basename(x[0]) == item]
                if not hit:
                    ap.error(f"--only {item}: no such file in {args.dir}")
                chosen += hit
        numbered = chosen
    lo, hi = args.from_, args.to or len(every)
    numbered = [x for x in numbered if lo <= x[1] <= hi]
    if args.first:
        numbered = numbered[:args.first]
    return numbered


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--dir", default="dlp_samples_high", help="sample folder: a name next to this script, or a path")
    ap.add_argument("--runs", type=int, default=1, help="times each file is sent (default 1)")
    ap.add_argument("--only", action="append", default=[],
                    help="only this file: its number (as # in the report) or name; repeatable")
    ap.add_argument("--first", "--limit", type=int, default=0, help="only the first N files (smallest first)")
    ap.add_argument("--from", dest="from_", type=int, default=1, help="start at file number N")
    ap.add_argument("--to", type=int, default=0, help="stop at file number N")
    ap.add_argument("--timeout", type=int, default=600, help="most seconds to wait for one response (default 600 = 10 min)")
    ap.add_argument("--stop-after-timeouts", type=int, default=2,
                    help="stop the run after this many timeouts in a row (default 2; 0 = never stop)")
    ap.add_argument("--retry-failed", action="store_true", help="re-send files whose earlier runs all failed")
    ap.add_argument("--mode", default="full", choices=["full", "fast"])
    ap.add_argument("--url", default=DEFAULT_URL)
    ap.add_argument("--token", default=os.environ.get("DLP_TOKEN"), help="bearer token, if the server requires one")
    ap.add_argument("--report", default=None, help="Excel report path (default: reports/<folder>_x<runs>.xlsx)")
    ap.add_argument("--log", default=None, help="jsonl of runs (default: next to the report)")
    ap.add_argument("--report-only", action="store_true", help="send nothing; rebuild the report from the log")
    args = ap.parse_args()

    if not os.path.isdir(args.dir) and os.path.isdir(os.path.join(HERE, args.dir)):
        args.dir = os.path.join(HERE, args.dir)
    if not os.path.isdir(args.dir):
        ap.error(f"--dir {args.dir}: folder not found")
    folder = os.path.basename(os.path.normpath(args.dir))
    args.report = args.report or os.path.join(HERE, "reports", f"{folder}_x{args.runs}.xlsx")
    os.makedirs(os.path.dirname(os.path.abspath(args.report)), exist_ok=True)
    log_path = args.log or os.path.splitext(args.report)[0] + ".jsonl"

    files = pick_files(ap, args)
    manifest = {}
    if os.path.exists(os.path.join(args.dir, "manifest.csv")):
        with open(os.path.join(args.dir, "manifest.csv"), encoding="utf-8") as fh:
            manifest = {r["file"]: r for r in csv.DictReader(fh)}
    runs_by_file = {}
    if os.path.exists(log_path):
        with open(log_path, encoding="utf-8") as fh:
            for line in fh:
                if line.strip():
                    r = json.loads(line)
                    runs_by_file.setdefault(r["file"], []).append(r)

    if args.report_only:
        files = [f for f in files if runs_by_file.get(os.path.basename(f[0]))]
        build_report(args.report, files, runs_by_file, manifest, args)
        print(f"Report for {len(files)} files: {args.report}", flush=True)
        return

    print(f"{folder}: {len(files)} file(s) x {args.runs} run(s) -> {args.url} (mode={args.mode}, "
          f"timeout {args.timeout}s, stop after {args.stop_after_timeouts} timeouts in a row)", flush=True)
    session = requests.Session()
    in_a_row = 0
    for path, number in files:
        name = os.path.basename(path)
        have = runs_by_file.setdefault(name, [])
        if args.retry_failed and have and not any(r.get("scan_ok") is True for r in have):
            print(f"[{number:02d}] {name:<34} retrying ({len(have)} failed run(s) dropped)", flush=True)
            have.clear()
            with open(log_path, "w", encoding="utf-8") as fh:
                for runs in runs_by_file.values():
                    fh.writelines(json.dumps(r) + "\n" for r in runs)
        if len(have) >= args.runs:
            continue
        planted = load_answer_key(args.dir, name)
        for run in range(len(have) + 1, args.runs + 1):
            r = run_once(session, args, path, args.timeout, planted)
            r["run"] = run
            have.append(r)
            with open(log_path, "a", encoding="utf-8") as fh:
                fh.write(json.dumps(r) + "\n")
            a = r.get("accuracy")
            status = ("OK" if r.get("scan_ok") is True else f"FAIL ({r.get('error') or r.get('http_status')})")
            acc = (f"  recall {a['recall']:.0%}  precision {a['precision']:.0%}  ({a['matched']}/{a['planted']} found, "
                   f"{a['extra']} extra)") if a and a["recall"] is not None and a["precision"] is not None else ""
            print(f"[{number:02d}] {name:<34} run {run}: {r['seconds']:>8.1f}s  {status}{acc}", flush=True)
            in_a_row = in_a_row + 1 if str(r.get("error", "")).startswith("Timed out") else 0
            if args.stop_after_timeouts and in_a_row >= args.stop_after_timeouts:
                break
        done = [f for f in files if runs_by_file.get(os.path.basename(f[0]))]
        build_report(args.report, done, runs_by_file, manifest, args)
        if args.stop_after_timeouts and in_a_row >= args.stop_after_timeouts:
            print(f"Stopped: {in_a_row} requests in a row hit the {args.timeout}s limit.", flush=True)
            break
    print(f"Report: {args.report}", flush=True)


if __name__ == "__main__":
    main()
